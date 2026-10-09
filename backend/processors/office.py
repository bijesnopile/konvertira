"""Spreadsheet and presentation conversions using direct APIs and LibreOffice."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from backend.formats import normalize_format, require_conversion
from backend.models.files import FileProcessingError, ProcessingErrorCode
from backend.processors.document import _libreoffice_convert
from config import settings


@dataclass(frozen=True, slots=True)
class OfficeResult:
    content: bytes
    extension: str
    content_type: str
    warning: str


def _safe_sheet_name(value: str, index: int) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", value).strip(".-")
    return (cleaned[:50] or f"sheet-{index}") + ".csv"


def _safe_spreadsheet_value(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _xlsx_to_csv(content: bytes) -> OfficeResult:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (InvalidFileException, KeyError, ValueError, OSError) as exc:
        raise FileProcessingError("The workbook is corrupt or unsupported.") from exc
    sheets: list[tuple[str, bytes]] = []
    try:
        for index, worksheet in enumerate(workbook.worksheets, start=1):
            text = io.StringIO(newline="")
            writer = csv.writer(text)
            for row in worksheet.iter_rows(values_only=True):
                writer.writerow(["" if value is None else _safe_spreadsheet_value(value) for value in row])
            sheets.append((_safe_sheet_name(worksheet.title, index), text.getvalue().encode("utf-8-sig")))
    finally:
        workbook.close()
    warning = "CSV does not preserve styles, charts, macros, hidden-sheet state, or workbook structure. Formula text is retained where available."
    if len(sheets) == 1:
        return OfficeResult(sheets[0][1], ".csv", "text/csv", warning)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in sheets:
            archive.writestr(name, data)
    return OfficeResult(output.getvalue(), ".zip", "application/zip", warning + " Each sheet is included as a separate generated CSV file.")


def _csv_to_xlsx(content: bytes, delimiter: str | None = None) -> bytes:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise FileProcessingError("CSV files must use UTF-8 encoding.") from exc
    if delimiter is not None and delimiter not in {",", ";", "\t", "|"}:
        raise FileProcessingError("CSV delimiter must be comma, semicolon, tab, or pipe.", 422)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|") if delimiter is None else None
    except csv.Error:
        dialect = csv.excel
    reader = (
        csv.reader(io.StringIO(text), delimiter=delimiter)
        if delimiter is not None
        else csv.reader(io.StringIO(text), dialect=dialect or "excel")
    )
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("Data")
    for row in reader:
        safe_row = [_safe_spreadsheet_value(value) for value in row]
        worksheet.append(safe_row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def convert_office(content: bytes, source_format: str, target_format: str, *, delimiter: str | None = None) -> OfficeResult:
    if not content:
        raise FileProcessingError("The file is empty.")
    if len(content) > settings.max_document_size:
        raise FileProcessingError("The file exceeds the size limit.", 413)
    source = normalize_format(source_format)
    target = normalize_format(target_format)
    if source is None or target is None:
        raise FileProcessingError("The requested format is unsupported.", 415)
    require_conversion(source.id, target.id, "server")

    warning = "Formulas, macros, charts, styles, transitions, fonts, embedded objects, and external links may change or be omitted. Macros are not deliberately executed."
    if target.id == "csv":
        workbook = content if source.id == "xlsx" else _libreoffice_convert(content, source.id, "xlsx")
        converted = _xlsx_to_csv(workbook)
    elif source.id == "csv":
        xlsx = _csv_to_xlsx(content, delimiter)
        if target.id == "xlsx":
            converted = OfficeResult(xlsx, ".xlsx", target.preferred_mime_type, "CSV has one table only. Formula-like text is escaped to reduce spreadsheet formula injection risk.")
        else:
            result = _libreoffice_convert(xlsx, "xlsx", "ods")
            converted = OfficeResult(result, ".ods", target.preferred_mime_type, "CSV has one table only. Formula-like text is escaped to reduce spreadsheet formula injection risk.")
    else:
        result = _libreoffice_convert(content, source.id, target.id)
        converted = OfficeResult(result, target.preferred_extension, target.preferred_mime_type, warning)
    if not converted.content:
        raise FileProcessingError("The conversion produced an empty file.", 500, ProcessingErrorCode.CONVERSION_FAILED)
    if len(converted.content) > settings.max_document_output_size:
        raise FileProcessingError("The generated file exceeds the output size limit.", 413, ProcessingErrorCode.DECODED_CONTENT_TOO_LARGE)
    return converted
