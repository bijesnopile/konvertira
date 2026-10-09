"""Direct text conversions and isolated LibreOffice document conversion."""

from __future__ import annotations

import html
import io
import os
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from bs4 import BeautifulSoup
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from markdown import markdown

from backend.formats import normalize_format, require_conversion
from backend.models.files import FileProcessingError, ProcessingErrorCode
from config import settings


@dataclass(frozen=True, slots=True)
class DocumentResult:
    content: bytes
    format_id: str
    extension: str
    content_type: str
    fidelity_warning: str | None = None


def _decode_utf8(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise FileProcessingError("Text files must use UTF-8 encoding.") from exc


def _docx_to_text(content: bytes) -> bytes:
    try:
        document = Document(io.BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        return text.encode("utf-8")
    except (PackageNotFoundError, ValueError, KeyError) as exc:
        raise FileProcessingError("The DOCX file is corrupt or unsupported.") from exc


def _text_to_docx(text: str) -> bytes:
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _markdown_to_html(content: bytes) -> bytes:
    # Escaping first prevents raw HTML, scripts, event handlers, and remote tags
    # from becoming active content. Markdown links remain ordinary links.
    text = html.escape(_decode_utf8(content), quote=False)
    body = markdown(text, extensions=["extra", "sane_lists"])
    soup = BeautifulSoup(body, "html.parser")
    for element in soup.find_all(True):
        if element.name == "a":
            href = str(element.get("href") or "")
            scheme = urlparse(href).scheme.lower()
            safe_href = href if scheme in {"", "http", "https", "mailto"} else ""
            element.attrs = {"href": safe_href, "rel": "noopener noreferrer"} if safe_href else {}
        else:
            element.attrs = {}
    body = str(soup)
    return ("<!doctype html><html><head><meta charset=\"utf-8\"></head><body>" + body + "</body></html>").encode("utf-8")


def _html_to_text(content: bytes) -> bytes:
    soup = BeautifulSoup(_decode_utf8(content), "html.parser")
    for element in soup(["script", "style", "iframe", "object", "embed", "svg"]):
        element.decompose()
    return soup.get_text("\n", strip=True).encode("utf-8")


def _libreoffice_convert(content: bytes, source: str, target: str) -> bytes:
    filters = {
        "pdf": "pdf", "docx": "docx", "odt": "odt",
        "xlsx": "xlsx", "ods": "ods", "pptx": "pptx", "odp": "odp",
    }
    target_filter = filters.get(target)
    if target_filter is None:
        raise FileProcessingError("This office output format is not supported.", 422)
    try:
        with tempfile.TemporaryDirectory(prefix="konvertira-document-") as directory:
            root = Path(directory)
            input_path = root / f"input.{source}"
            output_dir = root / "output"
            profile_dir = root / "profile"
            output_dir.mkdir()
            profile_dir.mkdir()
            input_path.write_bytes(content)
            command = [
                settings.libreoffice_path,
                "--headless", "--nologo", "--nodefault", "--nolockcheck", "--norestore",
                f"-env:UserInstallation={profile_dir.as_uri()}",
                "--convert-to", target_filter,
                "--outdir", str(output_dir),
                str(input_path),
            ]
            environment = {**os.environ, "HOME": str(root), "SAL_DISABLE_OPENCL": "1"}
            subprocess.run(
                command,
                cwd=root,
                env=environment,
                stdin=subprocess.DEVNULL,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=settings.document_job_timeout_seconds,
                check=False,
            )
            output_path = output_dir / f"input.{target}"
            if not output_path.is_file():
                raise FileProcessingError(
                    "The document conversion engine could not generate the requested output.",
                    422,
                    ProcessingErrorCode.CONVERSION_FAILED,
                )
            result = output_path.read_bytes()
            if not result or len(result) > settings.max_document_output_size:
                raise FileProcessingError("The generated document exceeds the output limit.", 413)
            return result
    except FileNotFoundError as exc:
        raise FileProcessingError(
            "The document conversion engine is unavailable.",
            503,
            ProcessingErrorCode.UNSUPPORTED_FEATURE,
        ) from exc
    except subprocess.TimeoutExpired as exc:
        raise FileProcessingError(
            "Document conversion timed out.",
            504,
            ProcessingErrorCode.PROCESSING_TIMEOUT,
        ) from exc


def convert_document(content: bytes, source_format: str, target_format: str) -> DocumentResult:
    if not content:
        raise FileProcessingError("The document is empty.")
    if len(content) > settings.max_document_size:
        raise FileProcessingError("The document exceeds the size limit.", 413, ProcessingErrorCode.FILE_TOO_LARGE)
    source = normalize_format(source_format)
    target = normalize_format(target_format)
    if source is None or target is None:
        raise FileProcessingError("The document format is not supported.", 415, ProcessingErrorCode.UNSUPPORTED_FORMAT)
    require_conversion(source.id, target.id, "server")

    warning: str | None = None
    if (source.id, target.id) == ("docx", "txt"):
        result = _docx_to_text(content)
        warning = "Formatting, images, headers, and complex document structures are not included in plain text."
    elif source.id == "txt" and target.id == "docx":
        result = _text_to_docx(_decode_utf8(content))
    elif source.id == "markdown" and target.id == "html":
        result = _markdown_to_html(content)
        warning = "Raw HTML is escaped and active content is not preserved."
    elif source.id == "html" and target.id == "txt":
        result = _html_to_text(content)
        warning = "Scripts, styles, embeds, and formatting are discarded."
    elif source.id == "markdown" and target.id in {"docx", "pdf"}:
        docx_content = _text_to_docx(_decode_utf8(content))
        result = docx_content if target.id == "docx" else _libreoffice_convert(docx_content, "docx", "pdf")
        warning = "Markdown is converted as readable text; complex layout and extensions may not be preserved."
    elif source.id == "txt" and target.id == "pdf":
        result = _libreoffice_convert(_text_to_docx(_decode_utf8(content)), "docx", "pdf")
    else:
        result = _libreoffice_convert(content, source.id, target.id)
        warning = "Complex layout, fonts, macros, embedded objects, and external links may change or be omitted."

    if not result:
        raise FileProcessingError("The document conversion produced an empty file.", 500, ProcessingErrorCode.CONVERSION_FAILED)
    if len(result) > settings.max_document_output_size:
        raise FileProcessingError("The generated document exceeds the output limit.", 413, ProcessingErrorCode.DECODED_CONTENT_TOO_LARGE)

    return DocumentResult(
        content=result,
        format_id=target.id,
        extension=target.preferred_extension,
        content_type=target.preferred_mime_type,
        fidelity_warning=warning,
    )
