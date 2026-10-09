import io
import shutil
import subprocess

import pytest
from docx import Document

from backend.models.files import FileProcessingError, ProcessingErrorCode
from backend.processors import document as processor


def make_docx(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_to_text_and_non_ascii() -> None:
    result = processor.convert_document(make_docx("Živjo svijete"), "docx", "txt")
    assert result.content.decode("utf-8") == "Živjo svijete"


def test_text_to_docx() -> None:
    result = processor.convert_document("Prvi\nDrugi".encode(), "txt", "docx")
    document = Document(io.BytesIO(result.content))
    assert [paragraph.text for paragraph in document.paragraphs] == ["Prvi", "Drugi"]


def test_markdown_html_escapes_active_content() -> None:
    result = processor.convert_document(
        b"# Hello\n<script>alert(1)</script><img src=x onerror=alert(2)>",
        "markdown",
        "html",
    )
    text = result.content.decode()
    assert "<script>" not in text
    assert "<img " not in text
    assert "&lt;script&gt;" in text


def test_markdown_html_rejects_javascript_links() -> None:
    result = processor.convert_document(b"[click](javascript:alert(1))", "markdown", "html")
    assert "javascript:" not in result.content.decode()


def test_html_to_text_drops_active_elements() -> None:
    result = processor.convert_document(
        b"<h1>Safe</h1><script>secret()</script><style>x</style>",
        "html",
        "txt",
    )
    assert result.content.decode() == "Safe"


def test_unsupported_document_conversion_is_rejected() -> None:
    with pytest.raises(FileProcessingError) as error:
        processor.convert_document(b"hello", "txt", "html")
    assert error.value.code == ProcessingErrorCode.UNSUPPORTED_CONVERSION


def test_libreoffice_timeout_maps_to_safe_error(monkeypatch) -> None:
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("hidden command", 1)

    monkeypatch.setattr(processor.subprocess, "run", timeout)
    with pytest.raises(FileProcessingError, match="timed out") as error:
        processor._libreoffice_convert(make_docx("hello"), "docx", "pdf")
    assert error.value.code == ProcessingErrorCode.PROCESSING_TIMEOUT


@pytest.mark.skipif(shutil.which("libreoffice") is None, reason="LibreOffice is not installed locally")
def test_libreoffice_docx_to_pdf_when_available() -> None:
    result = processor.convert_document(make_docx("hello"), "docx", "pdf")
    assert result.content.startswith(b"%PDF")
