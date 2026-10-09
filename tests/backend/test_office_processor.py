import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from backend.models.files import FileProcessingError
from backend.processors.office import convert_office


def make_xlsx(multi: bool = False) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "First sheet"
    worksheet.append(["name", "value"])
    worksheet.append(["Živjo", 42])
    worksheet.append(["formula", "=2+2"])
    if multi:
        second = workbook.create_sheet("Second sheet")
        second.append(["other"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_xlsx_to_single_csv() -> None:
    result = convert_office(make_xlsx(), "xlsx", "csv")
    assert result.extension == ".csv"
    assert "Živjo,42" in result.content.decode("utf-8-sig")
    assert "formula,'=2+2" in result.content.decode("utf-8-sig")


def test_multisheet_xlsx_to_safe_csv_zip() -> None:
    result = convert_office(make_xlsx(multi=True), "xlsx", "csv")
    assert result.extension == ".zip"
    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        assert archive.namelist() == ["First-sheet.csv", "Second-sheet.csv"]
        assert all("/" not in name and "\\" not in name for name in archive.namelist())


def test_utf8_csv_to_xlsx_and_formula_injection_mitigation() -> None:
    result = convert_office("ime;vrijednost\nČaj;=2+2\n".encode(), "csv", "xlsx", delimiter=";")
    workbook = load_workbook(io.BytesIO(result.content), data_only=False)
    worksheet = workbook.active
    assert worksheet["A2"].value == "Čaj"
    assert worksheet["B2"].value == "'=2+2"


def test_unsupported_office_pair() -> None:
    with pytest.raises(FileProcessingError, match="not supported"):
        convert_office(make_xlsx(), "xlsx", "pptx")
